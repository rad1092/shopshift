"""Bounded CSV/XLSX ingestion and lossless repeated ERP export merge.

XLSX uses openpyxl in read-only mode with data_only=False so formulas cannot
silently become stale cached values. Only the first worksheet is imported.
See https://openpyxl.readthedocs.io/en/stable/tutorial.html#loading-from-a-file
"""

from __future__ import annotations

import copy
import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from pathlib import Path

from .model import Operation, Project, Resource
from .timeutil import parse_timestamp, working_windows

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNZIPPED_BYTES = 80 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 64
MAX_CELL_CHARS = 4096
MAX_DURATION_SECONDS = 31 * 86400
MAX_RESOURCES = 2000
MAX_CALENDAR_WINDOWS = 100_000
FIELDS = ("id", "job", "machine", "operator", "setup", "run", "predecessors", "release", "due", "label")
REQUIRED = ("id", "job", "machine", "run")
UNITS = {"seconds": 1, "minutes": 60, "hours": 3600}


class ImportProblem(ValueError):
    """A source or mapping error; the original project was not modified."""


@dataclass
class ImportResult:
    project: Project
    added: list[str]
    updated: list[str]
    retained: list[str]
    warnings: list[str]


class _SourceRow(dict):
    """A normal mapping with the original CSV/XLSX record position for errors."""

    def __init__(self, values, source_row: int):
        super().__init__(values)
        self.source_row = source_row


class _SourceCells(list):
    def __init__(self, values, source_row: int):
        super().__init__(values)
        self.source_row = source_row


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(timespec="microseconds" if value.microsecond else "seconds")
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, bool):
        raise ImportProblem("Boolean spreadsheet cells are not supported; export text or numbers")
    else:
        text = str(value).strip()
    if len(text) > MAX_CELL_CHARS:
        raise ImportProblem(f"Cell exceeds {MAX_CELL_CHARS} characters")
    if any(ord(char) < 32 and char not in "\t\n\r" for char in text):
        raise ImportProblem("Input contains an unsupported control character")
    return text


def _records(raw_rows, *, pad_short: bool = False) -> tuple[list[str], list[dict[str, str]]]:
    iterator = iter(raw_rows)
    try:
        headers = [_cell(value) for value in next(iterator)]
    except StopIteration:
        raise ImportProblem("The input is empty") from None
    # Excel's formatting can extend beyond the actual header cells.
    while pad_short and headers and not headers[-1]:
        headers.pop()
    if not headers or len(headers) > MAX_COLUMNS:
        raise ImportProblem(f"Input needs 1–{MAX_COLUMNS} header columns")
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ImportProblem("Header names must be nonempty and unique")
    rows = []
    for number, values in enumerate(iterator, start=2):
        if number > MAX_ROWS + 1:
            raise ImportProblem(f"Input exceeds {MAX_ROWS} rows (including blank rows)")
        number = getattr(values, "source_row", number)
        cells = [_cell(value) for value in values]
        if not any(cells):
            continue
        if any(cells[len(headers):]) or (not pad_short and len(cells) > len(headers)):
            raise ImportProblem(f"Row {number} has cells beyond the header columns")
        if len(cells) < len(headers):
            if not pad_short:
                raise ImportProblem(f"Row {number} has {len(cells)} cells; expected {len(headers)}")
            cells.extend([""] * (len(headers) - len(cells)))
        rows.append(_SourceRow(zip(headers, cells[:len(headers)], strict=True), number))
    if not rows:
        raise ImportProblem("The input has headers but no operation rows")
    return headers, rows


def read_table(path) -> tuple[list[str], list[dict[str, str]]]:
    """Read UTF-8 CSV or the first XLSX worksheet without evaluating formulas."""
    path = Path(path)
    try:
        with path.open("rb") as handle:
            data = handle.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise ImportProblem("Input exceeds the 20 MiB file limit")
        if path.suffix.lower() == ".csv":
            text = data.decode("utf-8-sig")
            # strict=True rejects truncated quoted records. Never sniff dialect:
            # repeated saved mappings should behave identically across imports.
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)

            def csv_rows():
                start = 1
                for values in reader:
                    yield _SourceCells(values, start)
                    start = reader.line_num + 1

            return _records(csv_rows())
        if path.suffix.lower() != ".xlsx":
            raise ImportProblem("Choose .csv (UTF-8, comma separated) or .xlsx")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(item.file_size for item in members) > MAX_UNZIPPED_BYTES:
                raise ImportProblem("Expanded XLSX exceeds the 80 MiB / 2000 member safety limit")
            if any(item.flag_bits & 1 for item in members):
                raise ImportProblem("Encrypted XLSX files are not supported")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.worksheets:
                raise ImportProblem("XLSX has no worksheets")
            sheet = workbook.worksheets[0]
            if (sheet.max_row and sheet.max_row > MAX_ROWS + 1) or (
                sheet.max_column and sheet.max_column > MAX_COLUMNS
            ):
                raise ImportProblem(f"XLSX exceeds {MAX_ROWS} rows or {MAX_COLUMNS} columns")
            # Do not trust the producer's dimensions to omit hidden extra rows.
            sheet.reset_dimensions()

            def values():
                for number, row in enumerate(sheet.iter_rows(), start=1):
                    if number > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                        raise ImportProblem(f"XLSX exceeds {MAX_ROWS} rows or {MAX_COLUMNS} columns")
                    result = []
                    for cell in row:
                        if cell.data_type in ("f", "e"):
                            raise ImportProblem(f"XLSX cell {cell.coordinate} is a formula or error; export values first")
                        result.append(cell.value)
                    yield result

            return _records(values(), pad_short=True)
        finally:
            workbook.close()
    except ImportProblem:
        raise
    except Exception as exc:
        raise ImportProblem(f"Could not read {path.name}: {exc}") from exc


def _duration(text: str, unit: str, field: str, row: int) -> int:
    if len(text) > 40 or not re.fullmatch(r"(?:0|[1-9]\d*)(?:\.\d+)?", text):
        raise ImportProblem(f"Row {row}: {field} must be a nonnegative decimal in {unit}")
    try:
        seconds = Fraction(Decimal(text)) * UNITS[unit]
        if seconds.denominator != 1:
            raise ImportProblem(f"Row {row}: {field} must resolve to whole seconds; no duration rounding")
        result = int(seconds)
    except (InvalidOperation, OverflowError) as exc:
        raise ImportProblem(f"Row {row}: invalid {field}") from exc
    if not 0 <= result <= MAX_DURATION_SECONDS:
        raise ImportProblem(f"Row {row}: {field} exceeds 31 days")
    return result


def import_operations(project: Project, path, mapping: dict[str, str], duration_unit: str = "minutes") -> ImportResult:
    """Atomically merge IDs, retaining absent rows and every scenario decision.

Fields absent from the mapping retain existing values. A mapped blank optional
cell deliberately clears that field. New IDs use documented model defaults.
Changed durations/resources may invalidate retained decisions: they are never
silently moved. Caller must show warnings and the independent validation report.
"""
    if duration_unit not in UNITS:
        raise ImportProblem("Duration unit must be seconds, minutes, or hours")
    if not isinstance(mapping, dict) or any(
        key not in (*FIELDS, "__duration_unit") or not isinstance(value, str)
        for key, value in mapping.items()
    ):
        raise ImportProblem("Mapping contains unsupported fields or non-text column names")
    headers, rows = read_table(path)
    selected = {key: value for key, value in mapping.items() if key in FIELDS and value}
    missing = [field for field in REQUIRED if field not in selected]
    if missing:
        raise ImportProblem("Map required fields: " + ", ".join(missing))
    if len(set(selected.values())) != len(selected):
        raise ImportProblem("Each source column can map to only one field")
    for field, header in selected.items():
        if header not in headers:
            raise ImportProblem(f"Mapped column {header!r} for {field} is missing")
    imported = {}
    for number, row in enumerate(rows, start=2):
        number = getattr(row, "source_row", number)
        values = {field: row[header].strip() for field, header in selected.items()}
        for field in REQUIRED:
            if not values[field]:
                raise ImportProblem(f"Row {number}: {field} is required")
        identity = values["id"]
        for field in ("id", "job", "machine", "operator"):
            if field in values and (len(values[field]) > 200 or any(c in values[field] for c in "\r\n\t")):
                raise ImportProblem(f"Row {number}: {field} must fit in a single line of 200 characters")
        if identity in imported:
            raise ImportProblem(f"Row {number}: duplicate operation ID {identity!r}")
        old = project.operations.get(identity)
        operation = copy.deepcopy(old) if old else Operation(identity, values["job"], values["machine"])
        for field in ("job", "machine", "operator", "label"):
            if field in values:
                setattr(operation, field, values[field])
        for field in ("setup", "run"):
            if field in values:
                setattr(operation, field, _duration(values[field] or "0", duration_unit, field, number))
        if not 0 < operation.duration <= MAX_DURATION_SECONDS:
            raise ImportProblem(f"Row {number}: setup + run must be positive and at most 31 days")
        if "predecessors" in values:
            predecessors = [value.strip() for value in values["predecessors"].split(";") if value.strip()]
            if len(predecessors) != len(set(predecessors)) or identity in predecessors:
                raise ImportProblem(f"Row {number}: duplicate or self predecessor for {identity!r}")
            operation.predecessors = predecessors
        for field in ("release", "due"):
            if field in values:
                try:
                    setattr(operation, field, parse_timestamp(values[field], project.timezone) if values[field] else None)
                except ValueError as exc:
                    raise ImportProblem(f"Row {number}: {field}: {exc}") from exc
        imported[identity] = operation
    if len(set(project.operations) | set(imported)) > MAX_ROWS:
        raise ImportProblem(f"Merged project exceeds {MAX_ROWS} operations")
    # Missing predecessors can be imported to support partial ERP extracts;
    # the validator blocks suggestions until all dependencies are available.
    merged = copy.deepcopy(project)
    added = [identity for identity in imported if identity not in project.operations]
    updated = [identity for identity, operation in imported.items()
               if identity in project.operations and operation != project.operations[identity]]
    retained = [identity for identity in project.operations if identity not in imported]
    merged.operations.update(imported)
    warnings = []
    planned_resources = {}
    for attribute in ("machines", "operators"):
        field = "machine" if attribute == "machines" else "operator"
        resources = getattr(merged, attribute)
        new_ids = sorted({getattr(op, field) for op in imported.values()} - set(resources) - {""})
        planned_resources[attribute] = new_ids
    total_resources = len(merged.machines) + len(merged.operators) + sum(map(len, planned_resources.values()))
    if total_resources > MAX_RESOURCES:
        raise ImportProblem(f"Merged project exceeds {MAX_RESOURCES:,} machine/operator resources")
    windows = []
    if any(planned_resources.values()):
        try:
            windows = working_windows(merged.horizon_start, merged.horizon_end, merged.timezone)
        except (ValueError, OverflowError) as exc:
            raise ImportProblem(f"Set a valid project horizon before importing resources: {exc}") from exc
        existing_windows = sum(len(resource.windows) for resource in (*merged.machines.values(), *merged.operators.values()))
        if existing_windows + len(windows) * sum(map(len, planned_resources.values())) > MAX_CALENDAR_WINDOWS:
            raise ImportProblem(f"Default calendars would exceed {MAX_CALENDAR_WINDOWS:,} total windows; use a shorter planning horizon")
    for attribute, new_ids in planned_resources.items():
        resources = getattr(merged, attribute)
        if new_ids:
            for identity in new_ids:
                resources[identity] = Resource(identity, copy.deepcopy(windows))
            warnings.append(f"New {attribute}: {', '.join(new_ids)}. Assumed weekdays 08:00–12:00 / 13:00–17:00 "
                            f"in {merged.timezone}; review availability before use.")
    if retained:
        warnings.append(f"Retained {len(retained)} omitted operations; no deletion inferred from this export.")
    if updated:
        warnings.append(f"Updated {len(updated)} operations. Existing placements and locks in every scenario were "
                        "preserved; review conflicts after changed durations, resources or dependencies.")
    missing_links = [(op.id, pred) for op in imported.values() for pred in op.predecessors
                     if pred not in merged.operations]
    if missing_links:
        warnings.append(f"{len(missing_links)} predecessor references are missing; scheduling is blocked until resolved.")
    saved = dict(selected, __duration_unit=duration_unit)
    merged.mappings[Path(path).name] = saved
    merged.mappings["__last__"] = dict(saved)
    if Path(path).suffix.lower() == ".xlsx":
        warnings.append("Imported only the first XLSX worksheet; formulas and error cells are rejected.")
    return ImportResult(merged, added, updated, retained, warnings)
